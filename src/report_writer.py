"""표준양식 예산분석보고서 docx 생성기."""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor


KOREAN_FONT = "맑은 고딕"


@dataclass
class ReportData:
    """보고서에 채워질 데이터. 본문 3섹션은 규칙기반 엔진이 채운다."""

    project_name: str = ""
    department: str = ""
    fiscal_year: str = ""
    overview: str = ""          # 사업개요 및 현황
    issues: str = ""            # 문제점
    analysis: str = ""          # 분석의견
    raw_source_excerpt: str = ""  # 원본 자료 발췌 (참고용)


def _set_korean_font(run, size_pt: int = 11, bold: bool = False):
    run.font.name = KOREAN_FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rFonts is None:
        from docx.oxml.ns import qn
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    from docx.oxml.ns import qn
    rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    rFonts.set(qn("w:ascii"), KOREAN_FONT)
    rFonts.set(qn("w:hAnsi"), KOREAN_FONT)


def _add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_korean_font(run, size_pt=18, bold=True)


def _add_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_korean_font(run, size_pt=13, bold=True)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _add_body(doc: Document, text: str):
    if not text.strip():
        text = "(추후 작성)"
    for line in text.strip().split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_korean_font(run, size_pt=11)


def _add_meta_table(doc: Document, data: ReportData):
    table = doc.add_table(rows=3, cols=2)
    table.style = "Light Grid Accent 1"
    cells_data = [
        ("사업명", data.project_name or "(미입력)"),
        ("소관부서", data.department or "(미입력)"),
        ("작성일자", datetime.now().strftime("%Y-%m-%d")),
    ]
    for row, (label, value) in zip(table.rows, cells_data):
        for cell, txt in zip(row.cells, (label, value)):
            cell.text = ""
            run = cell.paragraphs[0].add_run(txt)
            _set_korean_font(run, size_pt=11, bold=(cell is row.cells[0]))


def build_report(data: ReportData) -> bytes:
    """ReportData를 받아 docx 바이트를 반환."""
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    _add_title(doc, "예산분석보고서 (초안)")
    doc.add_paragraph()

    _add_meta_table(doc, data)
    doc.add_paragraph()

    _add_heading(doc, "1. 사업개요 및 현황")
    _add_body(doc, data.overview)
    doc.add_paragraph()

    _add_heading(doc, "2. 문제점")
    _add_body(doc, data.issues)
    doc.add_paragraph()

    _add_heading(doc, "3. 분석의견")
    _add_body(doc, data.analysis)
    doc.add_paragraph()

    _add_heading(doc, "4. 참고 - 원본 자료 발췌")
    _add_body(doc, data.raw_source_excerpt[:2000] if data.raw_source_excerpt else "")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
